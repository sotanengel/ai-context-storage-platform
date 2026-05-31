"""Silver converter: XLSX → CDM via openpyxl."""

from __future__ import annotations

import io
from typing import Any

from formaforge.models.silver import (
    CdmBlock,
    CdmDocument,
    CdmFrontmatter,
    CdmTableBlock,
    ConversionMethod,
)
from formaforge.silver.converters.base import BaseConverter


def _open_xlsx(data: bytes) -> Any:
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required for XLSX ingestion. Install with: uv add 'formaforge[binary]'"
        ) from exc
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)


class XlsxConverter(BaseConverter):
    source_format = "xlsx"

    def convert(self, raw: str, source_uri: str = "") -> CdmDocument:
        return self.convert_bytes(raw.encode("utf-8", errors="replace"), source_uri)

    def convert_bytes(self, raw: bytes, source_uri: str = "") -> CdmDocument:
        workbook = _open_xlsx(raw)
        blocks: list[CdmBlock] = []
        title = ""

        for ws in workbook.worksheets:
            if not title:
                title = ws.title or ""
            all_rows = list(ws.iter_rows(values_only=False))
            if len(all_rows) < 2:
                continue
            headers = [str(cell.value) if cell.value is not None else "" for cell in all_rows[0]]
            columns = [f"{h}:string" for h in headers]
            data_rows = [[cell.value for cell in row] for row in all_rows[1:]]
            blocks.append(CdmTableBlock(columns=columns, rows=data_rows))

        fm = CdmFrontmatter(
            source_format=self.source_format,
            source_uri=source_uri,
            structure_class="structured",
            conversion_method=ConversionMethod.DETERMINISTIC,
            conversion_confidence=1.0,
        )
        return CdmDocument(frontmatter=fm, title=title, blocks=blocks)
